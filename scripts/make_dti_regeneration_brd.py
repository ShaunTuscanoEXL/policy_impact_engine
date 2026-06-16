"""Generate the engine-aligned 'Dynamic Choice Regeneration on DTI Breach'
BRD as a .docx for upload + extraction.

This is the original BRD reframed so our DECLARATIVE rule engine can
represent its essence. The original is largely a PROCEDURAL spec (loop:
regenerate N choices, recompute EMI/APR). Our engine is IF conditions →
SET/REJECT/ADJUST/FLAG on a single loan's fields, so:

  - The regeneration LOOP + EMI/APR recomputation are application/service
    layer (called out as out-of-engine-scope below).
  - The TRIGGER, ACCEPTANCE CRITERIA, and DECISION OUTCOME are
    expressible as rules and are written here in extractable form using
    our real loan-record field names.

Field alignment (verified against app.services.field_registry):
  debt_to_income_ratio -> calculated_attributes.debt_to_income_ratio
  net_monthly_surplus  -> calculated_attributes.net_monthly_surplus
  desired_amount       -> desired_amount
  interest_rate        -> decision_context.interest_rate
  (apr, ndi_ratio are NOT in the loan schema — see notes in the doc)
"""
from pathlib import Path
from docx import Document

OUT = Path(__file__).resolve().parents[1] / "data" / "uploads" / "BRD_DTI_Regeneration_Aligned.docx"
OUT.parent.mkdir(parents=True, exist_ok=True)

BRD = """BUSINESS REQUIREMENTS DOCUMENT (BRD)
Dynamic Choice Regeneration on DTI Breach — Decision Rules (Engine-Aligned V1.0)

Context (non-normative — not a rule):
When all generated loan choices are declined due to DTI breaches, the decision service regenerates lower loan amounts, recomputes affordability, and re-checks each candidate against the rules below. The regeneration loop and the EMI/DTI/APR recomputation are performed by the service; this document specifies ONLY the decision rules the policy engine enforces on each regenerated candidate. Every rule below applies only while the application is in regeneration, indicated by the field repeat_dti_regeneration set to REGENERATING.

Decision Rules (each row is one rule):

RULE 1 — Acceptance. When repeat_dti_regeneration equals REGENERATING and debt_to_income_ratio is less than or equal to 0.40 and net_monthly_surplus is greater than 0 and desired_amount is less than or equal to 60000 and interest_rate is less than or equal to 0.2999, set decision_status to APPROVED_WITH_CONDITIONS.

RULE 2 — Amount cap reject. When repeat_dti_regeneration equals REGENERATING and desired_amount is greater than 60000, set decision_status to REJECTED.

RULE 3 — DTI cap reject. When repeat_dti_regeneration equals REGENERATING and debt_to_income_ratio is greater than 0.40, set decision_status to REJECTED.

RULE 4 — Affordability reject. When repeat_dti_regeneration equals REGENERATING and net_monthly_surplus is less than or equal to 0, set decision_status to REJECTED.

RULE 5 — Pricing cap reject. When repeat_dti_regeneration equals REGENERATING and interest_rate is greater than 0.2999, set decision_status to REJECTED.

Field & threshold reference (for the rules above):
- repeat_dti_regeneration: scope flag; equals REGENERATING only during dynamic choice regeneration.
- debt_to_income_ratio: policy DTI cap is 0.40.
- net_monthly_surplus: must be greater than 0 (the affordability floor; the BRD's "NDI must stay positive" requirement).
- desired_amount: policy amount cap is 60000.
- interest_rate: pricing cap is 0.2999 (29.99%).

Out of engine scope (handled by the decision/pricing service, do not extract as rules):
APR cap of 36.99% (APR is computed per choice, not a stored loan field); the 70% retained-NDI ratio (no ratio field exists — enforced here as net_monthly_surplus > 0); the regeneration amount-boundary math and choice generation loop.
"""

doc = Document()
for line in BRD.strip().split("\n"):
    p = line.strip()
    if not p:
        doc.add_paragraph("")
        continue
    # Heading heuristic: "N." or "N.N" prefixes
    head = p.split(" ", 1)[0]
    if head and head[0].isdigit() and ("." in head):
        level = 1 if head.count(".") == 1 and head.endswith(".") is False and len(head) <= 3 else 2
        # crude: "5." -> level 1, "5.1" -> level 2
        level = 2 if head.count(".") >= 1 and not head.endswith(".") and len(head) > 2 else 1
        doc.add_heading(p, level=min(level, 2))
    elif p.startswith("- "):
        doc.add_paragraph(p[2:], style="List Bullet")
    else:
        doc.add_paragraph(p)

doc.save(OUT)
print(f"Wrote {OUT} ({OUT.stat().st_size} bytes)")
