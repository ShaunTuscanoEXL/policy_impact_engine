"""Generate 5 sample BRD documents in DOCX format for testing the extraction pipeline.

Each BRD has a different structure and complexity level to test robustness:

1. BRD_Formal_Structured.docx - Classic formal BRD with numbered sections
2. BRD_Policy_Memo.docx - Short informal policy memo (no section numbers)
3. BRD_Tabular_Rules.docx - Rules defined primarily in tables
4. BRD_Multi_Change.docx - Multiple policy changes in one document
5. BRD_Regulatory_Update.docx - Regulatory compliance update (dense, technical)
"""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from pathlib import Path
import datetime


def set_style(doc):
    """Set base font for the document."""
    style = doc.styles["Normal"]
    font = style.font
    font.name = "Calibri"
    font.size = Pt(11)


def add_title_page(doc, title, subtitle="", date=None):
    """Add a simple title."""
    p = doc.add_heading(title, level=0)
    if subtitle:
        p = doc.add_paragraph(subtitle)
        p.style.font.size = Pt(14)
        p.style.font.color.rgb = RGBColor(100, 100, 100)
    if date:
        doc.add_paragraph(f"Date: {date}")
    doc.add_paragraph("")


# ═══════════════════════════════════════════════════════
# BRD 1: Formal Structured (Classic numbered sections)
# ═══════════════════════════════════════════════════════
def create_brd1():
    doc = Document()
    set_style(doc)
    add_title_page(doc, "BRD-2026-001: Personal Loan Bureau Score Threshold Revision",
                   "Business Requirement Document", "April 2026")

    doc.add_heading("1.0 Executive Summary", level=1)
    doc.add_paragraph(
        "This BRD proposes revising the minimum bureau score threshold for personal loan "
        "eligibility from the current 650 to 680, and introducing a tiered interest rate "
        "structure based on bureau score bands. The change is driven by increasing default "
        "rates in the 650-679 score segment which have risen 23% year-over-year."
    )

    doc.add_heading("2.0 Background", level=1)
    doc.add_paragraph(
        "Our current personal loan underwriting policy sets the minimum bureau (CIBIL) score "
        "at 650 for all applicants. Analysis of the 2024-2025 portfolio shows that borrowers "
        "in the 650-679 range have a 12.3% default rate compared to 4.1% for those scoring "
        "680 and above. The risk-adjusted return on this segment is negative after accounting "
        "for collection costs."
    )

    doc.add_heading("3.0 Current State", level=1)
    doc.add_paragraph("The current lending rules for personal loans are:")
    doc.add_paragraph("• Minimum bureau score: 650", style="List Bullet")
    doc.add_paragraph("• Interest rate: Flat 14.5% for all approved applicants", style="List Bullet")
    doc.add_paragraph("• Maximum loan amount: ₹15,00,000", style="List Bullet")
    doc.add_paragraph("• Maximum DTI ratio: 50%", style="List Bullet")
    doc.add_paragraph("• Minimum monthly income: ₹25,000", style="List Bullet")

    doc.add_heading("4.0 Proposed Changes", level=1)

    doc.add_heading("4.1 Bureau Score Minimum", level=2)
    doc.add_paragraph(
        "Raise the minimum bureau score from 650 to 680. Applications with a bureau score "
        "below 680 should be automatically rejected with reason code 'BUREAU_SCORE_BELOW_MINIMUM'."
    )

    doc.add_heading("4.2 Tiered Interest Rate", level=2)
    doc.add_paragraph("Implement the following interest rate tiers based on bureau score:")

    table = doc.add_table(rows=5, cols=3)
    table.style = "Table Grid"
    headers = ["Bureau Score Range", "Risk Band", "Interest Rate"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    data = [
        ("680 - 719", "Medium-High", "16.5%"),
        ("720 - 749", "Medium", "14.5%"),
        ("750 - 799", "Low-Medium", "12.5%"),
        ("800+", "Low", "11.0%"),
    ]
    for row_idx, (score, band, rate) in enumerate(data, 1):
        table.rows[row_idx].cells[0].text = score
        table.rows[row_idx].cells[1].text = band
        table.rows[row_idx].cells[2].text = rate

    doc.add_heading("4.3 DTI Cap Reduction", level=2)
    doc.add_paragraph(
        "Reduce the maximum DTI ratio from 50% to 45% for applicants with bureau score "
        "below 750. Applicants with bureau score 750 and above retain the 50% DTI cap."
    )

    doc.add_heading("4.4 High-Value Loan Restriction", level=2)
    doc.add_paragraph(
        "For loan amounts exceeding ₹10,00,000 (10 lakh), require a minimum bureau score of "
        "720 and minimum monthly income of ₹50,000. Applications not meeting both criteria "
        "should be capped at ₹10,00,000 eligible amount."
    )

    doc.add_heading("5.0 Expected Impact", level=1)
    doc.add_paragraph(
        "Based on backtesting against 12 months of application data:\n"
        "• Approval rate reduction: ~8% (from 62% to ~57%)\n"
        "• Expected default rate reduction: 35% in first year\n"
        "• Portfolio quality improvement: Average risk band shifts from Medium-High to Medium\n"
        "• Revenue impact: +₹2.3Cr net (higher rates on medium-risk, fewer defaults)"
    )

    doc.add_heading("6.0 Implementation Timeline", level=1)
    doc.add_paragraph("• Phase 1 (Week 1-2): System configuration and rule updates")
    doc.add_paragraph("• Phase 2 (Week 3): UAT testing with sample data")
    doc.add_paragraph("• Phase 3 (Week 4): Production deployment")

    doc.save(str(output_dir / "BRD_Formal_Structured.docx"))
    print("Created: BRD_Formal_Structured.docx")


# ═══════════════════════════════════════════════════════
# BRD 2: Short Policy Memo (informal, no section numbers)
# ═══════════════════════════════════════════════════════
def create_brd2():
    doc = Document()
    set_style(doc)
    add_title_page(doc, "Policy Update: Self-Employed Applicant Criteria",
                   "Internal Memo - Credit Policy Team", "March 2026")

    doc.add_paragraph(
        "Following the quarterly portfolio review, we are tightening lending criteria for "
        "self-employed applicants effective May 1, 2026. The self-employed segment has shown "
        "elevated delinquency rates (18.7% vs 6.2% for salaried) and inconsistent income "
        "verification remains a challenge."
    )

    doc.add_paragraph("")
    doc.add_heading("What's Changing", level=2)

    doc.add_paragraph(
        "1. Minimum bureau score for self-employed applicants increases from 680 to 720.\n\n"
        "2. Monthly income floor raised to ₹40,000 (from ₹25,000) for self-employed.\n\n"
        "3. Banking stability index must be at least 0.65 (currently no minimum enforced).\n\n"
        "4. Salary credit consistency over 6 months must be at least 70% (0.70).\n\n"
        "5. Maximum loan amount for self-employed capped at ₹8,00,000 regardless of "
        "income (currently ₹15,00,000 same as salaried).\n\n"
        "6. DTI ratio cap tightened to 35% for self-employed (from 50%).\n\n"
        "7. Applicants with more than 3 active loans should be auto-rejected.\n\n"
        "8. Cash deposits exceeding ₹5,00,000 in the last 6 months will trigger a "
        "manual review flag."
    )

    doc.add_paragraph("")
    doc.add_heading("Rationale", level=2)
    doc.add_paragraph(
        "Self-employed applicants make up 28% of our loan book but account for 47% of NPAs. "
        "The new criteria target the highest-risk sub-segments while preserving access for "
        "well-qualified self-employed borrowers. Backtesting shows these criteria would have "
        "prevented 62% of defaults in this segment over the past 12 months."
    )

    doc.add_paragraph("")
    doc.add_heading("Exceptions", level=2)
    doc.add_paragraph(
        "Self-employed professionals (doctors, lawyers, CAs) with professional license "
        "verification may be exempted from the banking stability requirement at the "
        "discretion of the credit manager."
    )

    doc.save(str(output_dir / "BRD_Policy_Memo.docx"))
    print("Created: BRD_Policy_Memo.docx")


# ═══════════════════════════════════════════════════════
# BRD 3: Tabular Rules (rules defined in tables)
# ═══════════════════════════════════════════════════════
def create_brd3():
    doc = Document()
    set_style(doc)
    add_title_page(doc, "BRD-2026-003: Credit Risk Scoring Model Update",
                   "Risk Analytics Division", "February 2026")

    doc.add_heading("Purpose", level=1)
    doc.add_paragraph(
        "Update the credit risk scoring thresholds and introduce new risk bands to align "
        "with the revised regulatory guidelines from RBI circular 2025/14. This impacts "
        "eligibility, pricing, and exposure limits across all unsecured lending products."
    )

    doc.add_heading("New Risk Band Definitions", level=1)
    doc.add_paragraph("The following risk bands replace the existing 3-tier system with a 5-tier model:")

    table = doc.add_table(rows=6, cols=5)
    table.style = "Table Grid"
    h = ["Risk Band", "Bureau Score", "Max DTI", "Interest Rate", "Max Exposure"]
    for i, val in enumerate(h):
        table.rows[0].cells[i].text = val
    rows = [
        ("Prime", "800+", "50%", "10.5%", "₹25,00,000"),
        ("Near-Prime", "750-799", "45%", "12.5%", "₹15,00,000"),
        ("Standard", "700-749", "40%", "14.5%", "₹10,00,000"),
        ("Sub-Standard", "680-699", "35%", "17.0%", "₹5,00,000"),
        ("Decline", "Below 680", "N/A", "N/A", "₹0 (Auto-reject)"),
    ]
    for r, (band, score, dti, rate, exp) in enumerate(rows, 1):
        table.rows[r].cells[0].text = band
        table.rows[r].cells[1].text = score
        table.rows[r].cells[2].text = dti
        table.rows[r].cells[3].text = rate
        table.rows[r].cells[4].text = exp

    doc.add_heading("Additional Eligibility Gates", level=1)
    doc.add_paragraph("The following hard stops apply before risk band assignment:")

    table2 = doc.add_table(rows=7, cols=4)
    table2.style = "Table Grid"
    h2 = ["Rule #", "Condition", "Threshold", "Action"]
    for i, val in enumerate(h2):
        table2.rows[0].cells[i].text = val
    gates = [
        ("G-01", "Max DPD in last 12 months", "> 30 days", "Auto-reject"),
        ("G-02", "Overdue accounts", "> 0", "Auto-reject"),
        ("G-03", "Credit inquiries in last 3 months", "> 5", "Flag for manual review"),
        ("G-04", "Cheque bounces in last 6 months", "> 2", "Auto-reject"),
        ("G-05", "Unsecured loans count", "> 4", "Flag for manual review"),
        ("G-06", "Credit utilization ratio", "> 80%", "Reduce max exposure by 30%"),
    ]
    for r, (rule_num, cond, threshold, action) in enumerate(gates, 1):
        table2.rows[r].cells[0].text = rule_num
        table2.rows[r].cells[1].text = cond
        table2.rows[r].cells[2].text = threshold
        table2.rows[r].cells[3].text = action

    doc.add_heading("G5/G6 Score Integration", level=1)
    doc.add_paragraph(
        "Effective immediately, the internal G5 score must be above 450 and G6 score "
        "above 500 for all applications. Applications failing either threshold should "
        "be flagged for senior credit officer review regardless of bureau score."
    )

    doc.add_heading("Minimum Income Requirements", level=1)
    table3 = doc.add_table(rows=4, cols=3)
    table3.style = "Table Grid"
    h3 = ["City Tier", "Min Monthly Income", "Min Account Vintage"]
    for i, val in enumerate(h3):
        table3.rows[0].cells[i].text = val
    income_rows = [
        ("Tier 1", "₹30,000", "6 months"),
        ("Tier 2", "₹25,000", "12 months"),
        ("Tier 3", "₹20,000", "18 months"),
    ]
    for r, (tier, income, vintage) in enumerate(income_rows, 1):
        table3.rows[r].cells[0].text = tier
        table3.rows[r].cells[1].text = income
        table3.rows[r].cells[2].text = vintage

    doc.save(str(output_dir / "BRD_Tabular_Rules.docx"))
    print("Created: BRD_Tabular_Rules.docx")


# ═══════════════════════════════════════════════════════
# BRD 4: Multiple Policy Changes (complex, multi-rule)
# ═══════════════════════════════════════════════════════
def create_brd4():
    doc = Document()
    set_style(doc)
    add_title_page(doc, "Quarterly Policy Revision - Q2 2026",
                   "Credit Policy Committee Decision Record", "April 2026")

    doc.add_paragraph(
        "This document captures all lending policy changes approved in the Q2 2026 "
        "Credit Policy Committee meeting held on April 5, 2026. Changes are effective "
        "from April 15, 2026 unless otherwise noted."
    )

    doc.add_heading("Change 1: Young Borrower Premium", level=1)
    doc.add_paragraph(
        "Applicants under 25 years of age will attract a 1.5% interest rate premium "
        "on top of their risk-band rate. This applies only to first-time borrowers "
        "(repeat_type = 'NEW'). The rationale is the higher delinquency observed in "
        "this age group (14.2% vs portfolio average 7.8%).\n\n"
        "Conditions: age < 25 AND repeat_type == 'NEW'\n"
        "Action: interest_rate += 1.5%"
    )

    doc.add_heading("Change 2: Repeat Borrower Benefit", level=1)
    doc.add_paragraph(
        "Existing customers with at least one successfully closed loan and no overdue "
        "history will receive a 0.5% interest rate discount and a 20% increase in "
        "maximum eligible amount.\n\n"
        "Conditions: repeat_type == 'REPEAT' AND closed_loans >= 1 AND overdue_accounts == 0\n"
        "Actions:\n"
        "  - interest_rate -= 0.5%\n"
        "  - eligible_amount *= 1.20 (20% increase)"
    )

    doc.add_heading("Change 3: High DTI Warning", level=1)
    doc.add_paragraph(
        "Applications where the debt-to-income ratio exceeds 40% but is below the "
        "rejection threshold should be flagged for manual credit review rather than "
        "auto-approved. Currently these are auto-approved if other criteria are met.\n\n"
        "Condition: dti_ratio > 0.40 AND dti_ratio <= 0.50\n"
        "Action: Flag for manual review (do not auto-approve)"
    )

    doc.add_heading("Change 4: City Tier 3 Restrictions", level=1)
    doc.add_paragraph(
        "For applications from Tier 3 cities, apply the following additional restrictions:\n"
        "• Maximum loan amount capped at ₹5,00,000\n"
        "• Minimum employment tenure increased to 24 months (from 12)\n"
        "• Minimum bureau score increased to 720 (from 680)\n\n"
        "These restrictions aim to address the 22% higher default rate observed in "
        "Tier 3 city applications."
    )

    doc.add_heading("Change 5: Digital Lending Scorecard", level=1)
    doc.add_paragraph(
        "Introduce a net monthly surplus floor of ₹10,000 for all applications. "
        "Applicants whose calculated net monthly surplus (income minus all obligations) "
        "falls below ₹10,000 should be rejected regardless of other criteria.\n\n"
        "Additionally, applicants with a transaction volatility index above 0.75 should "
        "be flagged for enhanced due diligence."
    )

    doc.add_heading("Change 6: EMI Bounce Protection", level=1)
    doc.add_paragraph(
        "Applicants with more than 2 loan repayment bounces in the last 12 months "
        "should be auto-rejected. Those with exactly 1-2 bounces should have their "
        "eligible amount reduced by 25%.\n\n"
        "Condition 1: loan_repayment_bounces_12m > 2 → REJECT\n"
        "Condition 2: loan_repayment_bounces_12m >= 1 AND loan_repayment_bounces_12m <= 2 "
        "→ eligible_amount *= 0.75"
    )

    doc.save(str(output_dir / "BRD_Multi_Change.docx"))
    print("Created: BRD_Multi_Change.docx")


# ═══════════════════════════════════════════════════════
# BRD 5: Regulatory Compliance Update (dense, freeform)
# ═══════════════════════════════════════════════════════
def create_brd5():
    doc = Document()
    set_style(doc)
    add_title_page(doc, "Regulatory Compliance Update - RBI Fair Lending Guidelines 2026",
                   "Compliance & Legal Team", "April 2026")

    doc.add_paragraph(
        "In compliance with the Reserve Bank of India's updated Fair Lending Practice "
        "Guidelines (Circular No. RBI/2026/45, dated March 15, 2026), the following "
        "modifications to our lending criteria are mandated effective June 1, 2026."
    )

    doc.add_paragraph("")
    doc.add_paragraph(
        "The circular requires all NBFCs to implement the following safeguards in their "
        "personal loan underwriting process. Non-compliance attracts penalties under "
        "Section 45-IA of the RBI Act."
    )

    doc.add_heading("Income Verification Mandate", level=2)
    doc.add_paragraph(
        "All applicants must have a verified monthly income of at least ₹15,000. "
        "The system must reject applications where monthly_income < 15000 with "
        "reason code 'RBI_MIN_INCOME_NOT_MET'. This is a hard regulatory floor "
        "and cannot be overridden by any approval authority."
    )

    doc.add_heading("Maximum Leverage Ratio", level=2)
    doc.add_paragraph(
        "The total loan exposure (desired_amount) must not exceed 10 times the "
        "applicant's monthly income. For example, an applicant earning ₹50,000 "
        "per month cannot be approved for more than ₹5,00,000. If desired_amount > "
        "monthly_income * 10, the eligible amount must be capped at monthly_income * 10."
    )

    doc.add_heading("Credit Bureau Mandate", level=2)
    doc.add_paragraph(
        "All applicants must have a valid bureau score. Applications with bureau_score "
        "of 0 or missing bureau data must be declined. The minimum acceptable bureau "
        "score is 650 for regulatory purposes (individual lenders may set higher "
        "thresholds). Applications with bureau_score < 650 must be rejected with "
        "reason 'RBI_BUREAU_MINIMUM'."
    )

    doc.add_heading("Cooling Period After Rejection", level=2)
    doc.add_paragraph(
        "Applicants with more than 6 credit inquiries in the last 3 months indicate "
        "possible credit-hungry behavior. Such applications (inquiries_last_3m > 6) "
        "must be auto-declined. Additionally, applicants with inquiries_last_3m between "
        "4 and 6 must be flagged for enhanced due diligence."
    )

    doc.add_heading("Unsecured Exposure Limit", level=2)
    doc.add_paragraph(
        "To prevent over-leveraging in unsecured debt, applicants with more than "
        "5 active unsecured loans must be declined. Those with 4-5 unsecured loans "
        "should have their eligible amount reduced by 40%. Applicants with 3 unsecured "
        "loans get a 20% reduction in eligible amount."
    )

    doc.add_heading("Interest Rate Ceiling", level=2)
    doc.add_paragraph(
        "Per the revised usury guidelines, the maximum interest rate for personal "
        "loans cannot exceed 24.0% per annum. Any risk-based pricing that would "
        "result in a rate above 24.0% must be capped at 24.0%. The system should "
        "flag such cases for compliance review."
    )

    doc.add_heading("Minimum Age Requirement", level=2)
    doc.add_paragraph(
        "Applicants must be at least 21 years of age for unsecured personal loans. "
        "Applications with age < 21 must be rejected with reason 'RBI_MIN_AGE'. "
        "Additionally, applicants above 60 years must provide additional income "
        "verification — flag these for manual document review."
    )

    doc.save(str(output_dir / "BRD_Regulatory_Update.docx"))
    print("Created: BRD_Regulatory_Update.docx")


# ═══════════════════════════════════════════════════════
# Main
# ═══════════════════════════════════════════════════════
output_dir = Path(__file__).parent
output_dir.mkdir(parents=True, exist_ok=True)

if __name__ == "__main__":
    print("Generating 5 sample BRDs...")
    create_brd1()
    create_brd2()
    create_brd3()
    create_brd4()
    create_brd5()
    print(f"\nAll BRDs saved to: {output_dir}")
